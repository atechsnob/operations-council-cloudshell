import os
import subprocess
import sys

# Ensure python-docx is installed
try:
    from docx import Document
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    from docx import Document

files = [
    "chapter_01_setup/README.md",
    "chapter_02_agent_studio/README.md",
    "chapter_03_council/README.md",
    "chapter_04_memory_and_mcp/README.md",
    "chapter_05_trial_by_fire/README.md",
    "chapter_06_governance/README.md",
    "chapter_07_deployment/README.md",
    "chapter_08_wrap_up/README.md"
]

doc = Document()
doc.add_heading('AI Service Desk: Architecting a Multi-Agent Support Team', level=1)

for f in files:
    if not os.path.exists(f):
        continue
    
    with open(f, "r") as inf:
        content = inf.read().replace("Operations Council", "AI Service Desk")
        
    lines = content.split('\n')
    in_code_block = False
    code_text = []
    
    for line in lines:
        if line.startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_text = []
                doc.add_paragraph("Code block")
                doc.add_paragraph("")
            else:
                in_code_block = False
                if code_text:
                    p = doc.add_paragraph('\n'.join(code_text))
                doc.add_paragraph("")
            continue
            
        if in_code_block:
            code_text.append(line)
            continue
            
        if line.startswith("# "):
            # [Heading 2] per the template instructions for top-level headers
            doc.add_heading(line[2:], level=2)
            # Also insert the literal text annotation just in case they rely on search/replace
            p = doc.add_paragraph(f"[Heading 2] {line[2:]}")
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=3)
            p = doc.add_paragraph(f"[Heading 3] {line[3:]}")
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=4)
        elif line.startswith("- ") or line.startswith("* "):
            doc.add_paragraph(line[2:], style='List Bullet')
        elif line.startswith("---"):
            pass
        elif line.strip() == "":
            pass
        elif line.startswith("> "):
            doc.add_paragraph(line[2:])
        else:
            doc.add_paragraph(line)

    doc.add_page_break()

doc.save("LIL_Article_Manuscript.docx")
print("Saved to LIL_Article_Manuscript.docx")
